import os
import sys
import tempfile
import pytest
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from pipeline.docx_exporter import DocxExporter


class TestDocxExporter:

    def test_markdown_to_docx_content(self):
        temp_dir = tempfile.mkdtemp()
        md_content = """# Título de Apuntes

## 1. Resumen
Este es un párrafo de **prueba** con texto en *cursiva* y `código`.

- Elemento de lista 1
- Elemento de lista 2

```python
def hello():
    print("Mundo")
```
"""
        md_file = os.path.join(temp_dir, "test_notes.md")
        with open(md_file, "w", encoding="utf-8") as f:
            f.write(md_content)

        docx_path = DocxExporter.convert_file(md_file)
        assert os.path.exists(docx_path)
        assert docx_path.endswith(".docx")

        doc = Document(docx_path)
        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
        assert "Título de Apuntes" in paragraphs_text
        assert "1. Resumen" in paragraphs_text
        assert any("prueba" in p for p in paragraphs_text)

    def test_markdown_to_docx_with_table(self):
        temp_dir = tempfile.mkdtemp()
        md_content = """# Tabla Test
| Encabezado 1 | Encabezado 2 |
| :--- | :--- |
| Celda A1 | Celda A2 <br>• Línea 2 |
| Celda B1 | Celda B2 |
"""
        md_file = os.path.join(temp_dir, "test_table.md")
        with open(md_file, "w", encoding="utf-8") as f:
            f.write(md_content)

        docx_path = DocxExporter.convert_file(md_file)
        assert os.path.exists(docx_path)
        doc = Document(docx_path)
        assert len(doc.tables) == 1
        assert len(doc.tables[0].rows) == 3
        assert len(doc.tables[0].columns) == 2
        assert "Encabezado 1" in doc.tables[0].cell(0, 0).text
        assert "Celda A1" in doc.tables[0].cell(1, 0).text
        assert "Línea 2" in doc.tables[0].cell(1, 1).text
